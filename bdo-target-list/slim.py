"""Collapse the working workbook into the two-tab deliverable Lance asked for:
Target List, and Excluded with the reason for every row that did not make it.
"""
import sys, re
sys.path.insert(0,'.')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from lib import norm_name, norm_domain
from collections import Counter

src_wb = openpyxl.load_workbook("BDO Target List - Combined Send File.xlsx")

def category(reason):
    """Reasons arrive prefixed ("Not a buy-box business: Above buy-box"), so look for the
    distinctive marker anywhere in the string and check the specific cases before the generic."""
    r = (reason or "").lower()
    if "no longer independent" in r or "acquired by" in r:   return "Acquired - no longer independent"
    if "pe-owned" in r or "pe platform" in r or "pe-backed" in r or "owned by" in r:
        return "PE-owned or subsidiary"
    if "above buy-box" in r:                                  return "Above buy-box"
    if "badlands-owned" in r:                                 return "Badlands-owned"
    if "broker" in r or "advisor" in r:                       return "Broker or advisor"
    if "duplicate" in r:                                      return "Duplicate record"
    if "guard" in r:                                          return "Guard services"
    return "Not a buy-box business"

rows=[]
seen=set()
def push(company, web, city, state, vert, cat, why):
    key=(norm_name(company or ""), (state or "").upper())
    if key in seen: return
    seen.add(key)
    if cat == "Acquired - no longer independent":
        web = ""   # the domain now resolves to the buyer, so it is not this company's address
    rows.append([company, web, city, state, vert, cat, re.sub(r"\s+"," ",(why or "")).strip()])

# 2) live deals - not excluded on the merits, held back so BDO is not told talks exist
for r in src_wb["Hold Back - Live Deals"].iter_rows(min_row=2, values_only=True):
    push(r[0], r[1], r[2], r[3], r[4],
         "Live deal - held back for your call",
         f"In the M&A pipeline at {r[5]}. Held back so BDO is not told talks exist; release it if you want the intro.")

# 1) screened out - acquisitions first, so they are never filed under a softer heading
_removed_rows=[(r, category(r[5] or "")) for r in src_wb["Removed"].iter_rows(min_row=2, values_only=True)]
ACQ_BY_NAME={norm_name(k):v for k,v in {
    "Day Automation":"Acquired by Stark Tech; its site says 'Day Automation is now Stark Tech'",
    "ENE Systems":"Acquired by Stark Tech; enesystems.com now resolves to starktech.com",
    "Advanced Door Service":"Acquired by Door Services Corporation; domain resolves to doorservicescorporation.com",
}.items()}
for r,c in _removed_rows:
    if c=="Acquired - no longer independent": ACQ_BY_NAME[norm_name(r[0] or "")]=r[5]
for want_acq in (True, False):
    for r,cat in _removed_rows:
        if (cat=="Acquired - no longer independent") != want_acq: continue
        reason=r[5] or ""
        why=re.sub(r"^(Not a buy-box business|Excluded \(NYC screen\)|Excluded \(universe screen\)|No longer independent):\s*","",reason).strip()
        push(r[0], r[1], r[2], r[3], r[4], cat, why or reason)

# 3) adjacent
for r in src_wb["Adjacent - Review"].iter_rows(min_row=2, values_only=True):
    _acq=ACQ_BY_NAME.get(norm_name(r[0] or ""))
    if _acq:
        push(r[0], r[1], r[2], r[3], r[4], "Acquired - no longer independent",
             re.sub(r"^No longer independent:\s*","",_acq))
    else:
        push(r[0], r[1], r[2], r[3], r[4], "Adjacent - your call",
             r[7] or "Distributor, wholesale monitoring, or an electrical/controls contractor with a security line. Not core buy-box.")

# 4) NYC unclassified
for r in src_wb["NYC Unclassified - Review"].iter_rows(min_row=2, values_only=True):
    ind = r[4] if r[4] and r[4]!="NO_MATCH" else ""
    emp = f", {r[5]} employees" if r[5] else ""
    bits=[b for b in [ind+emp if ind else "", r[8] or ""] if b]
    push(r[0], r[1], r[2], r[3], "",
         "NYC unclassified - your call",
         ("; ".join(bits)+". " if bits else "")+
         "Left unclassified by the September NYC cleanup; ZoomInfo's industry code does not separate low-voltage installers from IT shops here.")

# 5) the blocklists themselves - companies that never reached the candidate pool because
# they were screened out upstream. Included so this file answers "why isn't X here?" on its own.
from lib import read, D
for f, cat, default in [
    ("univ__peowned_and_subsidiaries.csv","PE-owned or subsidiary","Owned by a PE platform or operating as a subsidiary."),
    ("univ__above_buybox.csv","Above buy-box","Too large for a Badlands deal. Useful as BDO relationship context, not as a target."),
    ("univ__excluded.csv","Not a buy-box business","Outside the buy box."),
]:
    for x in read(D+f):
        why = (x.get("Ownership Note") or "").strip() if cat=="PE-owned or subsidiary" else ""
        why = why or (x.get("Exclusion Reason") or "").strip() or (x.get("Notes") or "").strip() or default
        push(x.get("Company"), x.get("Website"), x.get("HQ City"), x.get("State"),
             x.get("Vertical"), cat, why)
for x in read(D+"nyc__excluded.csv"):
    push(x.get("Company Name"), "", x.get("City"), x.get("State"), x.get("Vertical"),
         "Guard services" if "guard" in (x.get("Exclusion Reason") or "").lower() else "Not a buy-box business",
         x.get("Exclusion Reason") or "Screened out of the NYC metro list.")

# 6) broker and advisor companies attached to deals - removed so BDO is not handed
# Badlands' intermediaries as if they were targets.
for _b in ["Flatirons Capital Advisors","Hedgestone Advisors","DGP Capital","Corum Group",
           "The Advisory Investment Bank","SEMM Holdings"]:
    push(_b, "", "", "", "", "Broker or advisor",
         "Investment bank or advisor attached to a deal in the pipeline, not an operating target.")

# A duplicate record is merged into the row that survived, not excluded on its merits.
rows=[r for r in rows if r[5]!="Duplicate record"]
# Nothing may appear on both tabs: the Target List is authoritative.
_on_target={(norm_name(r[0] or ""), (r[3] or "").upper())
            for r in src_wb["BDO Send List"].iter_rows(min_row=2, values_only=True)}
_before=len(rows)
rows=[r for r in rows if (norm_name(r[0] or ""), (r[3] or "").upper()) not in _on_target]
print(f"dropped from Excluded: duplicates, plus {_before-len(rows)} rows that survived onto the Target List")

ORDER=["Live deal - held back for your call","Acquired - no longer independent","PE-owned or subsidiary",
       "Above buy-box","Adjacent - your call","NYC unclassified - your call","Guard services",
       "Not a buy-box business","Broker or advisor","Badlands-owned"]
rows.sort(key=lambda x:(ORDER.index(x[5]) if x[5] in ORDER else 99, (x[0] or "").lower()))

# ---- write the two-tab book ----
out=openpyxl.Workbook(); out.remove(out.active)
HDR=Font(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F3864")
def sheet(title, headers, data, widths):
    ws=out.create_sheet(title)
    ws.append(headers)
    for c in ws[1]: c.font=HDR; c.fill=FILL; c.alignment=Alignment(vertical="center")
    for row in data: ws.append(row)
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{len(data)+1}"
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    return ws

send=[list(r) for r in src_wb["BDO Send List"].iter_rows(min_row=2, values_only=True)]
sheet("Target List", ["Company","Website","City","State","Vertical","Size Tier","In HubSpot"],
      send, [40,32,20,7,32,13,12])
ws=sheet("Excluded", ["Company","Website","City","State","Vertical","Why excluded","Detail"],
      rows, [40,30,18,7,28,34,95])
for r in range(2, ws.max_row+1):
    ws.cell(r,7).alignment=Alignment(wrap_text=True, vertical="top")

out.save("BDO Target List.xlsx")
print("Target List:",len(send)," Excluded:",len(rows))
for k,v in Counter(r[5] for r in rows).most_common(): print(f"   {v:5d}  {k}")
